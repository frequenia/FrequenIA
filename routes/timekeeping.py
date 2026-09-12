"""Consultas e exportações de marcações para a gestão.

Todas as saídas são derivadas de ``marcacoes``. Os instantes continuam como
``timestamptz``; a conversão ocorre somente para filtro e apresentação.
"""

from collections import defaultdict
import csv
from datetime import date, datetime, time, timedelta, timezone
import io
from uuid import UUID
from zoneinfo import ZoneInfo

import psycopg2.extras
from docx import Document
from flask import Blueprint, g, jsonify, request, send_file
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from db import conectar_bd
from utils.auth_decorator import require_roles

timekeeping_bp = Blueprint("timekeeping", __name__)
MANAGEMENT_ROLES = ("administrador", "gestor", "rh")
LOCAL_TIMEZONE = ZoneInfo("America/Sao_Paulo")
CLOCK_TYPES = ("entrada", "saida_intervalo", "retorno_intervalo", "saida")


def _parse_uuid(value, field_name):
    try:
        return str(UUID(str(value)))
    except (TypeError, ValueError, AttributeError) as exc:
        raise ValueError(f"{field_name} inválido.") from exc


def _parse_date(value, field_name):
    if not value:
        return None
    try:
        return date.fromisoformat(str(value))
    except (TypeError, ValueError) as exc:
        raise ValueError(f"{field_name} inválida.") from exc


def _parse_filters(args):
    if args.get("empresa_id") not in (None, ""):
        raise ValueError("A consulta não aceita empresa informada pelo cliente.")
    employee_id = _parse_uuid(args.get("funcionario_id"), "Funcionário")
    start_date = _parse_date(args.get("inicio"), "Data inicial")
    end_date = _parse_date(args.get("fim"), "Data final")
    if start_date and end_date and end_date < start_date:
        raise ValueError("A data final não pode anteceder a inicial.")
    return employee_id, start_date, end_date


def _utc_bounds(start_date, end_date):
    """Converte datas civis de São Paulo em limites UTC indexáveis."""
    start = None
    end_exclusive = None
    if start_date:
        start = datetime.combine(start_date, time.min, LOCAL_TIMEZONE).astimezone(
            timezone.utc
        )
    if end_date:
        end_exclusive = datetime.combine(
            end_date + timedelta(days=1), time.min, LOCAL_TIMEZONE
        ).astimezone(timezone.utc)
    return start, end_exclusive


def _format_duration(seconds):
    if seconds is None or seconds < 0:
        return "--"
    total_minutes = int(seconds // 60)
    hours, minutes = divmod(total_minutes, 60)
    return f"{hours}h{minutes:02d}"


def _daily_total(events):
    entrada = events.get("entrada")
    saida = events.get("saida")
    if not entrada or not saida or saida < entrada:
        return "--"

    saida_intervalo = events.get("saida_intervalo")
    retorno_intervalo = events.get("retorno_intervalo")
    if bool(saida_intervalo) != bool(retorno_intervalo):
        return "--"

    total_seconds = (saida - entrada).total_seconds()
    if saida_intervalo and retorno_intervalo:
        if not (entrada <= saida_intervalo <= retorno_intervalo <= saida):
            return "--"
        total_seconds -= (retorno_intervalo - saida_intervalo).total_seconds()
    return _format_duration(total_seconds)


def _aware_instant(value):
    # psycopg2 retorna timestamptz consciente de fuso. Este fallback deixa
    # mocks ingênuos explicitamente UTC, sem usar o relógio do dispositivo.
    return value if value.tzinfo else value.replace(tzinfo=timezone.utc)


def _serialize_daily_records(rows):
    grouped = defaultdict(dict)
    weekdays = ["Segunda", "Terça", "Quarta", "Quinta", "Sexta", "Sábado", "Domingo"]
    ordered_rows = sorted(
        rows,
        key=lambda item: (_aware_instant(item["instante"]), str(item.get("id", ""))),
    )
    for row in ordered_rows:
        if row["tipo"] not in CLOCK_TYPES:
            continue
        local_instant = _aware_instant(row["instante"]).astimezone(LOCAL_TIMEZONE)
        local_date = local_instant.date()
        # Duplicatas: primeira marcação confirmada de cada tipo no dia local.
        # Nenhum registro histórico é apagado ou alterado.
        grouped[local_date].setdefault(row["tipo"], local_instant)

    result = []
    for local_date in sorted(grouped):
        events = grouped[local_date]
        formatted = {
            event_type: (
                events[event_type].strftime("%H:%M")
                if events.get(event_type)
                else "--:--"
            )
            for event_type in CLOCK_TYPES
        }
        result.append(
            {
                "data": local_date.isoformat(),
                "dia": weekdays[local_date.weekday()],
                "entrada": formatted["entrada"],
                "saida_intervalo": formatted["saida_intervalo"],
                "retorno_intervalo": formatted["retorno_intervalo"],
                "volta_intervalo": formatted["retorno_intervalo"],
                "saida": formatted["saida"],
                "total": _daily_total(events),
            }
        )
    return result


def _fetch_employee_records(cursor, company_id, employee_id, start_date, end_date):
    cursor.execute(
        """
        SELECT f.id AS funcionario_id, u.nome, f.matricula
        FROM funcionarios f
        INNER JOIN usuarios u ON u.id = f.usuario_id
        WHERE f.id = %s AND f.empresa_id = %s
        """,
        (employee_id, company_id),
    )
    employee = cursor.fetchone()
    if not employee:
        return None, []

    start_utc, end_utc_exclusive = _utc_bounds(start_date, end_date)
    query = """
        SELECT id, tipo, instante
        FROM marcacoes
        WHERE empresa_id = %s
          AND funcionario_id = %s
          AND estado = 'confirmada'
    """
    params = [company_id, employee_id]
    if start_utc:
        query += " AND instante >= %s"
        params.append(start_utc)
    if end_utc_exclusive:
        query += " AND instante < %s"
        params.append(end_utc_exclusive)
    query += " ORDER BY instante ASC, id ASC"
    cursor.execute(query, tuple(params))
    return employee, _serialize_daily_records(cursor.fetchall())


def _load_management_records(company_id, employee_id, start_date, end_date):
    connection = conectar_bd()
    cursor = connection.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        return _fetch_employee_records(
            cursor, company_id, employee_id, start_date, end_date
        )
    finally:
        cursor.close()
        connection.close()


def _format_date_br(value):
    return date.fromisoformat(value).strftime("%d/%m/%Y")


def _export_csv(records):
    output = io.StringIO()
    writer = csv.writer(output, delimiter=";")
    writer.writerow(
        ["Dia", "Data", "Entrada", "Saída Int.", "Retorno Int.", "Saída", "Total"]
    )
    for item in records:
        writer.writerow(
            [
                item["dia"],
                _format_date_br(item["data"]),
                item["entrada"],
                item["saida_intervalo"],
                item["retorno_intervalo"],
                item["saida"],
                item["total"],
            ]
        )
    payload = io.BytesIO(output.getvalue().encode("utf-8-sig"))
    payload.seek(0)
    return payload


def _draw_pdf_header(pdf, y):
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(40, y, "Relatório de Controle de Ponto")
    y -= 30
    pdf.setFont("Helvetica", 9)
    positions = (40, 95, 160, 220, 295, 370, 430)
    labels = ("Dia", "Data", "Entrada", "Saída Int.", "Retorno Int.", "Saída", "Total")
    for x, label in zip(positions, labels):
        pdf.drawString(x, y, label)
    y -= 15
    pdf.line(40, y, 550, y)
    return y


def _export_pdf(records):
    payload = io.BytesIO()
    pdf = canvas.Canvas(payload, pagesize=A4)
    _, page_height = A4
    y = _draw_pdf_header(pdf, page_height - 40)
    for item in records:
        y -= 18
        if y < 50:
            pdf.showPage()
            y = _draw_pdf_header(pdf, page_height - 40) - 18
        values = (
            item["dia"],
            _format_date_br(item["data"]),
            item["entrada"],
            item["saida_intervalo"],
            item["retorno_intervalo"],
            item["saida"],
            item["total"],
        )
        for x, value in zip((40, 95, 160, 220, 295, 370, 430), values):
            pdf.drawString(x, y, str(value))
    pdf.save()
    payload.seek(0)
    return payload


def _export_docx(records):
    document = Document()
    document.add_heading("Relatório de Controle de Ponto", level=1)
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ("Dia", "Data", "Entrada", "Saída Int.", "Retorno Int.", "Saída", "Total")
    for cell, label in zip(table.rows[0].cells, headers):
        cell.text = label
    for item in records:
        values = (
            item["dia"],
            _format_date_br(item["data"]),
            item["entrada"],
            item["saida_intervalo"],
            item["retorno_intervalo"],
            item["saida"],
            item["total"],
        )
        for cell, value in zip(table.add_row().cells, values):
            cell.text = str(value)
    payload = io.BytesIO()
    document.save(payload)
    payload.seek(0)
    return payload


@timekeeping_bp.get("/api/gestao/funcionarios")
@require_roles(*MANAGEMENT_ROLES)
def funcionarios_gestao():
    connection = conectar_bd()
    cursor = connection.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        cursor.execute(
            """
            SELECT f.id AS funcionario_id, u.nome, f.matricula
            FROM funcionarios f
            INNER JOIN usuarios u ON u.id = f.usuario_id
            WHERE f.empresa_id = %s
              AND f.status = 'ativo'
              AND u.status = 'ativo'
            ORDER BY u.nome, f.id
            """,
            (g.auth_context["empresa_id"],),
        )
        return jsonify(
            [
                {
                    "funcionario_id": str(item["funcionario_id"]),
                    "nome": item["nome"],
                    "matricula": item["matricula"],
                }
                for item in cursor.fetchall()
            ]
        )
    except Exception:
        return jsonify({"erro": "Não foi possível carregar os funcionários."}), 500
    finally:
        cursor.close()
        connection.close()


@timekeeping_bp.get("/api/gestao/pontos")
@require_roles(*MANAGEMENT_ROLES)
def pontos_gestao():
    try:
        employee_id, start_date, end_date = _parse_filters(request.args)
        employee, records = _load_management_records(
            g.auth_context["empresa_id"], employee_id, start_date, end_date
        )
        if not employee:
            return jsonify({"erro": "Funcionário não encontrado."}), 404
        return jsonify(records)
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400
    except Exception:
        return jsonify({"erro": "Não foi possível consultar as marcações."}), 500


@timekeeping_bp.get("/exportar-pontos")
@require_roles(*MANAGEMENT_ROLES)
def exportar_pontos_gestao():
    try:
        export_format = str(request.args.get("formato") or "").lower()
        if export_format not in {"csv", "pdf", "word", "doc", "docx"}:
            return jsonify({"erro": "Formato inválido."}), 400
        employee_id, start_date, end_date = _parse_filters(request.args)
        employee, records = _load_management_records(
            g.auth_context["empresa_id"], employee_id, start_date, end_date
        )
        if not employee:
            return jsonify({"erro": "Funcionário não encontrado."}), 404
        if not records:
            return jsonify({"erro": "Nenhum registro encontrado para exportação."}), 404

        if export_format == "csv":
            return send_file(
                _export_csv(records),
                as_attachment=True,
                download_name="controle_ponto.csv",
                mimetype="text/csv",
            )
        if export_format == "pdf":
            return send_file(
                _export_pdf(records),
                as_attachment=True,
                download_name="controle_ponto.pdf",
                mimetype="application/pdf",
            )
        return send_file(
            _export_docx(records),
            as_attachment=True,
            download_name="controle_ponto.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400
    except Exception:
        return jsonify({"erro": "Não foi possível exportar as marcações."}), 500
