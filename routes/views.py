from flask import (
    Blueprint,
    app,
    current_app,
    g,
    render_template,
    redirect,
    url_for,
    Flask,
    jsonify,
    session,
    request,
    send_file,
)
from datetime import datetime, date, timedelta, timezone
from decimal import Decimal, InvalidOperation
import hashlib
import re
from uuid import UUID
from zoneinfo import ZoneInfo
from utils.auth_decorator import access_token_required, login_required, require_roles
from utils.schedules import (
    parse_iso_date,
    parse_optional_iso_date,
    serialize_period,
    validate_periods,
    validate_timezone,
)
import psycopg2.extras
import jwt
import secrets
from werkzeug.security import generate_password_hash, check_password_hash
from db import (
    buscar_usuario_login_por_cpf,
    buscar_vinculos_ativos,
    conectar_bd,
)
from routes.face import pasta_usuario
from collections import defaultdict
import os
import csv
import io

from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from docx import Document

views_bp = Blueprint("views", __name__)

JWT_ALGORITHM = "HS256"
REFRESH_COOKIE_NAME = "frequenia_refresh_token"
ALLOWED_EMPLOYEE_PROFILES = {"funcionario", "gestor", "rh", "administrador"}
ALLOWED_CONTRACT_TYPES = {
    "efetivo",
    "comissionado",
    "temporario",
    "estagiario",
    "terceirizado",
    "outro",
}
EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def normalizar_cpf(value):
    return re.sub(r"\D", "", value or "")


def cpf_valido(value):
    cpf = normalizar_cpf(value)
    if len(cpf) != 11 or cpf == cpf[0] * 11:
        return False

    numbers = [int(digit) for digit in cpf]
    first_sum = sum(numbers[index] * (10 - index) for index in range(9))
    first_remainder = first_sum % 11
    first_digit = 0 if first_remainder < 2 else 11 - first_remainder
    second_sum = sum(numbers[index] * (11 - index) for index in range(10))
    second_remainder = second_sum % 11
    second_digit = 0 if second_remainder < 2 else 11 - second_remainder
    return first_digit == numbers[9] and second_digit == numbers[10]


def uuid_obrigatorio(value, field_name):
    try:
        return str(UUID(str(value)))
    except (TypeError, ValueError, AttributeError) as exc:
        raise ValueError(f"{field_name} inválido.") from exc


def empresa_id_autorizada(value):
    empresa_id = uuid_obrigatorio(value, "Empresa")
    if empresa_id != str(g.auth_context["empresa_id"]):
        raise PermissionError("Empresa fora do escopo autorizado.")
    return empresa_id


def empresa_autenticada(dados=None):
    empresa_id = str(g.auth_context["empresa_id"])
    if dados and dados.get("empresa_id") not in (None, ""):
        empresa_id_autorizada(dados["empresa_id"])
    return empresa_id


def validar_turno_payload(dados, partial=False):
    if not isinstance(dados, dict):
        raise ValueError("Corpo da requisição inválido.")

    result = {}
    if not partial or "nome" in dados:
        nome = str(dados.get("nome") or "").strip()
        if not nome:
            raise ValueError("Nome do turno é obrigatório.")
        result["nome"] = nome

    if not partial or "timezone" in dados:
        result["timezone"] = validate_timezone(dados.get("timezone"))

    if not partial or "status" in dados:
        status = str(dados.get("status") or "ativo").strip().lower()
        if status not in {"ativo", "inativo"}:
            raise ValueError("Status do turno inválido.")
        result["status"] = status

    if not partial or "periodos" in dados:
        result["periodos"] = validate_periods(dados.get("periodos"))

    if partial and not result:
        raise ValueError("Nenhuma alteração de turno foi informada.")
    return result


def inserir_periodos_turno(cursor, empresa_id, turno_id, periodos):
    for periodo in periodos:
        cursor.execute(
            """
            INSERT INTO periodos_turno (
                empresa_id, turno_id, dia_semana, ordem,
                inicio, fim, fim_dia_offset
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (
                empresa_id,
                turno_id,
                periodo["dia_semana"],
                periodo["ordem"],
                periodo["inicio"],
                periodo["fim"],
                periodo["fim_dia_offset"],
            ),
        )


def buscar_turno(cursor, empresa_id, turno_id, bloquear=False):
    cursor.execute(
        f"""
        SELECT id, nome, timezone, status
        FROM turnos
        WHERE id = %s AND empresa_id = %s
        {"FOR UPDATE" if bloquear else ""}
        """,
        (turno_id, empresa_id),
    )
    turno = cursor.fetchone()
    if not turno:
        return None

    cursor.execute(
        """
        SELECT id, dia_semana, ordem, inicio, fim, fim_dia_offset
        FROM periodos_turno
        WHERE turno_id = %s AND empresa_id = %s
        ORDER BY dia_semana, ordem, id
        """,
        (turno_id, empresa_id),
    )
    turno["periodos"] = [serialize_period(item) for item in cursor.fetchall()]
    return turno


def serializar_turno(turno):
    return {
        "id": str(turno["id"]),
        "nome": turno["nome"],
        "timezone": turno["timezone"],
        "status": turno["status"],
        "periodos": turno["periodos"],
    }


def buscar_jornada_data(cursor, empresa_id, funcionario_id, data_consulta):
    cursor.execute(
        """
        SELECT
            ft.id AS atribuicao_id,
            ft.vigencia_inicio,
            ft.vigencia_fim,
            t.id AS turno_id,
            t.nome AS turno_nome,
            t.timezone,
            t.status AS turno_status
        FROM funcionarios_turnos ft
        INNER JOIN turnos t
            ON t.id = ft.turno_id
           AND t.empresa_id = ft.empresa_id
        WHERE ft.empresa_id = %s
          AND ft.funcionario_id = %s
          AND ft.vigencia_inicio <= %s
          AND (ft.vigencia_fim IS NULL OR ft.vigencia_fim >= %s)
        ORDER BY ft.vigencia_inicio DESC, ft.created_at DESC, ft.id
        LIMIT 1
        """,
        (empresa_id, funcionario_id, data_consulta, data_consulta),
    )
    atribuicao = cursor.fetchone()
    if not atribuicao:
        return None

    cursor.execute(
        """
        SELECT id, dia_semana, ordem, inicio, fim, fim_dia_offset
        FROM periodos_turno
        WHERE empresa_id = %s AND turno_id = %s
        ORDER BY dia_semana, ordem, id
        """,
        (empresa_id, atribuicao["turno_id"]),
    )
    return {
        "turno": {
            "id": str(atribuicao["turno_id"]),
            "nome": atribuicao["turno_nome"],
            "timezone": atribuicao["timezone"],
            "status": atribuicao["turno_status"],
        },
        "vigencia": {
            "inicio": atribuicao["vigencia_inicio"].isoformat(),
            "fim": (
                atribuicao["vigencia_fim"].isoformat()
                if atribuicao["vigencia_fim"]
                else None
            ),
        },
        "periodos": [serialize_period(item) for item in cursor.fetchall()],
    }


def validar_dados_administrativos(dados, password_required=False):
    nome = str(dados.get("nome") or "").strip()
    email = str(dados.get("email") or "").strip().lower()
    telefone = str(dados.get("telefone") or "").strip()
    cpf = normalizar_cpf(dados.get("cpf"))
    matricula = str(dados.get("matricula") or "").strip()
    perfil = str(dados.get("perfil") or "").strip()
    tipo_contrato = str(dados.get("tipo_contrato") or "").strip()
    data_admissao_value = str(dados.get("data_admissao") or "").strip()

    if not nome:
        raise ValueError("Nome é obrigatório.")
    if not EMAIL_PATTERN.fullmatch(email):
        raise ValueError("E-mail inválido.")
    if not cpf_valido(cpf):
        raise ValueError("CPF inválido.")
    if not matricula:
        raise ValueError("Matrícula é obrigatória.")
    if perfil not in ALLOWED_EMPLOYEE_PROFILES:
        raise ValueError("Perfil inválido.")
    if tipo_contrato not in ALLOWED_CONTRACT_TYPES:
        raise ValueError("Tipo de contrato inválido.")

    try:
        data_admissao = date.fromisoformat(data_admissao_value)
    except ValueError as exc:
        raise ValueError("Data de admissão inválida.") from exc

    try:
        carga_horaria = Decimal(str(dados.get("carga_horaria_semanal") or ""))
    except InvalidOperation as exc:
        raise ValueError("Carga horária inválida.") from exc
    if carga_horaria <= 0 or carga_horaria > 168:
        raise ValueError("Carga horária inválida.")

    password = str(dados.get("senha") or "")
    if password_required and len(password) < 12:
        raise ValueError("A senha deve possuir pelo menos 12 caracteres.")

    return {
        "nome": nome,
        "email": email,
        "telefone": telefone or None,
        "cpf": cpf,
        "matricula": matricula,
        "perfil": perfil,
        "tipo_contrato": tipo_contrato,
        "data_admissao": data_admissao,
        "carga_horaria_semanal": carga_horaria,
        "empresa_id": uuid_obrigatorio(dados.get("empresa_id"), "Empresa"),
        "unidade_id": uuid_obrigatorio(dados.get("unidade_id"), "Unidade"),
        "equipe_id": uuid_obrigatorio(dados.get("equipe_id"), "Equipe"),
        "cargo_id": uuid_obrigatorio(dados.get("cargo_id"), "Cargo"),
        "senha": password,
    }


def validar_estrutura_funcional(cursor, dados):
    cursor.execute(
        """
        SELECT 1
        FROM empresas e
        INNER JOIN unidades un
            ON un.empresa_id = e.id AND un.id = %s AND un.status = 'ativa'
        INNER JOIN equipes eq
            ON eq.empresa_id = e.id
           AND eq.unidade_id = un.id
           AND eq.id = %s
           AND eq.status = 'ativa'
        INNER JOIN cargos c
            ON c.empresa_id = e.id AND c.id = %s AND c.status = 'ativo'
        WHERE e.id = %s AND e.status = 'ativa'
        """,
        (
            dados["unidade_id"],
            dados["equipe_id"],
            dados["cargo_id"],
            dados["empresa_id"],
        ),
    )
    if not cursor.fetchone():
        raise ValueError("Empresa, unidade, equipe ou cargo incompatível.")


def gerar_access_token(user_id, funcionario_id, session_id):
    agora = datetime.now(timezone.utc)
    payload = {
        "sub": str(user_id),
        "funcionario_id": str(funcionario_id),
        "sid": str(session_id),
        "type": "access",
        "iat": agora,
        "exp": agora
        + timedelta(minutes=current_app.config["JWT_ACCESS_TOKEN_MINUTES"]),
    }
    return jwt.encode(
        payload,
        current_app.config["JWT_SECRET_KEY"],
        algorithm=JWT_ALGORITHM,
    )


def gerar_refresh_token():
    raw_token = secrets.token_urlsafe(48)
    token_hash = hashlib.sha256(raw_token.encode("utf-8")).hexdigest()
    return raw_token, token_hash


def criar_auth_session(cursor, user_id, vinculo, familia_id=None):
    refresh_token, refresh_token_hash = gerar_refresh_token()
    cursor.execute(
        """
        INSERT INTO auth_sessions (
            usuario_id, empresa_id, funcionario_id,
            refresh_token_hash, familia_id, expires_at
        )
        VALUES (
            %s, %s, %s, %s, COALESCE(%s, extensions.gen_random_uuid()),
            now() + (%s * interval '1 day')
        )
        RETURNING id, familia_id
        """,
        (
            user_id,
            vinculo["empresa_id"],
            vinculo["funcionario_id"],
            refresh_token_hash,
            familia_id,
            current_app.config["JWT_REFRESH_TOKEN_DAYS"],
        ),
    )
    auth_session = cursor.fetchone()
    return auth_session, refresh_token


def resposta_com_tokens(payload, refresh_token, status=200):
    response = jsonify(payload)
    response.set_cookie(
        REFRESH_COOKIE_NAME,
        refresh_token,
        max_age=current_app.config["JWT_REFRESH_TOKEN_DAYS"] * 24 * 60 * 60,
        httponly=True,
        secure=current_app.config["REFRESH_COOKIE_SECURE"],
        samesite="Lax",
        path="/auth",
    )
    return response, status


def limpar_cookie_refresh(response):
    response.delete_cookie(
        REFRESH_COOKIE_NAME,
        httponly=True,
        secure=current_app.config["REFRESH_COOKIE_SECURE"],
        samesite="Lax",
        path="/auth",
    )
    return response


def revogar_familia_por_sessao(session_id):
    conn = conectar_bd()
    cursor = conn.cursor()
    try:
        cursor.execute(
            """
            UPDATE auth_sessions
            SET revoked_at = COALESCE(revoked_at, now())
            WHERE familia_id = (
                SELECT familia_id FROM auth_sessions WHERE id = %s
            )
            """,
            (session_id,),
        )
        conn.commit()
        return cursor.rowcount > 0
    except Exception:
        conn.rollback()
        raise
    finally:
        cursor.close()
        conn.close()


@views_bp.route("/")
def home():
    return redirect(url_for("views.login_page"))


# ==================================================================================================
# RENDERIZAÇÃO - PÁGINAS DO SISTEMA
# ==================================================================================================
@views_bp.route("/controleponto")
@login_required
def controle_ponto():
    usuarios = []

    if session.get("tipo") == "admin":
        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cursor.execute("""
            SELECT id, nome
            FROM usuarios
            ORDER BY nome
        """)

        usuarios = cursor.fetchall()

        cursor.close()
        conn.close()

    return render_template("controleponto.html", usuarios=usuarios)


@views_bp.route("/inicio")
def inicio():
    return render_template("inicio.html")


@views_bp.route("/cadastroUsuario")
@require_roles("administrador")
def cadastro_usuario():
    return render_template("cadastroUsuario.html")


@views_bp.route("/reconhecimentoFacial")
def reconhecimento_facial():
    return render_template("reconhecimentoFacial.html")


@views_bp.route("/cadastrarFoto")
@login_required
def cadastrar_foto():
    return render_template("cadastrarFoto.html")


@views_bp.route("/redefinicaoSenha")
def redefinicao_senha():
    return render_template("redefinicaoSenha.html")


@views_bp.route("/gerenciarUsuario")
@require_roles("administrador")
def gerenciar_usuario():
    return render_template("gerenciarUsuario.html")


@views_bp.route("/jornadas")
@require_roles("administrador")
def jornadas():
    return render_template("jornadas.html")


@views_bp.route("/gerenciarEmpresa")
@require_roles("administrador")
def gerenciar_empresa():
    return render_template("gerenciarEmpresa.html")


@views_bp.route("/configuracoes")
@require_roles("administrador")
def configuracoes():
    return render_template("configuracoes.html")


@views_bp.route("/cadastroEmpresas")
@require_roles("administrador")
def cadastroEmpresas():
    return render_template("cadastroEmpresas.html")


@views_bp.route("/recuperacaoSenha")
def recuperacao_senha():
    return render_template("recuperacaoSenha.html")


@views_bp.route("/inserirToken")
def inserir_token():
    return render_template("inserirToken.html")


# ==================================================================================================
# FUNÇÃO PRINCIPAL - LOGIN
# ==================================================================================================
@views_bp.route("/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}

    cpf = data.get("cpf")
    senha = data.get("senha")

    if not isinstance(cpf, str) or not isinstance(senha, str):
        return jsonify({"erro": "Credenciais inválidas."}), 401

    cpf = cpf.replace(".", "").replace("-", "").strip()
    senha = senha.strip()

    if not cpf or not senha:
        return jsonify({"erro": "Credenciais inválidas."}), 401

    user = buscar_usuario_login_por_cpf(cpf)

    if (
        not user
        or user["status"] != "ativo"
        or not user["senha_hash"]
        or not check_password_hash(user["senha_hash"], senha)
    ):
        return jsonify({"erro": "Credenciais inválidas."}), 401

    vinculos = buscar_vinculos_ativos(user["id"])
    if not vinculos:
        return jsonify({"erro": "Credenciais inválidas."}), 401

    if len(vinculos) > 1:
        return jsonify({"erro": "Seleção de vínculo necessária."}), 409

    vinculo = vinculos[0]
    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        auth_session, refresh_token = criar_auth_session(cursor, user["id"], vinculo)
        conn.commit()
    except Exception:
        conn.rollback()
        return jsonify({"erro": "Não foi possível iniciar a sessão."}), 500
    finally:
        cursor.close()
        conn.close()

    session["user_id"] = str(user["id"])
    session["nome"] = user["nome"]
    session["tipo"] = vinculo["perfil"]
    session["funcionario_id"] = str(vinculo["funcionario_id"])
    session["empresa_id"] = str(vinculo["empresa_id"])
    session["auth_session_id"] = str(auth_session["id"])
    access_token = gerar_access_token(
        user["id"],
        vinculo["funcionario_id"],
        auth_session["id"],
    )

    return resposta_com_tokens(
        {
            "ok": True,
            "nome": user["nome"],
            "tipo": vinculo["perfil"],
            "funcionario_id": str(vinculo["funcionario_id"]),
            "access_token": access_token,
            "refresh_token": refresh_token,
            "token_type": "Bearer",
            "expires_in": current_app.config["JWT_ACCESS_TOKEN_MINUTES"] * 60,
        },
        refresh_token,
    )


@views_bp.post("/auth/refresh")
def auth_refresh():
    data = request.get_json(silent=True) or {}
    refresh_token = data.get("refresh_token") or request.cookies.get(
        REFRESH_COOKIE_NAME
    )
    if not isinstance(refresh_token, str) or not refresh_token:
        return jsonify({"erro": "Refresh token inválido."}), 401

    refresh_token_hash = hashlib.sha256(refresh_token.encode("utf-8")).hexdigest()
    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        cursor.execute(
            """
            SELECT
                id, usuario_id, empresa_id, funcionario_id, familia_id,
                expires_at, rotated_at, revoked_at
            FROM auth_sessions
            WHERE refresh_token_hash = %s
            FOR UPDATE
            """,
            (refresh_token_hash,),
        )
        current_session = cursor.fetchone()
        if not current_session:
            conn.rollback()
            response = jsonify({"erro": "Refresh token inválido."})
            return limpar_cookie_refresh(response), 401

        if current_session["rotated_at"] is not None:
            cursor.execute(
                """
                UPDATE auth_sessions
                SET revoked_at = COALESCE(revoked_at, now())
                WHERE familia_id = %s
                """,
                (current_session["familia_id"],),
            )
            conn.commit()
            response = jsonify({"erro": "Refresh token inválido."})
            return limpar_cookie_refresh(response), 401

        if (
            current_session["revoked_at"] is not None
            or current_session["expires_at"] <= datetime.now(timezone.utc)
        ):
            conn.rollback()
            response = jsonify({"erro": "Refresh token inválido."})
            return limpar_cookie_refresh(response), 401

        cursor.execute(
            """
            SELECT 1
            FROM usuarios u
            INNER JOIN funcionarios f
                ON f.usuario_id = u.id
               AND f.id = %s
               AND f.empresa_id = %s
            INNER JOIN empresas e ON e.id = f.empresa_id
            WHERE u.id = %s
              AND u.status = 'ativo'
              AND f.status = 'ativo'
              AND e.status = 'ativa'
            """,
            (
                current_session["funcionario_id"],
                current_session["empresa_id"],
                current_session["usuario_id"],
            ),
        )
        if not cursor.fetchone():
            cursor.execute(
                """
                UPDATE auth_sessions
                SET revoked_at = COALESCE(revoked_at, now())
                WHERE familia_id = %s
                """,
                (current_session["familia_id"],),
            )
            conn.commit()
            response = jsonify({"erro": "Refresh token inválido."})
            return limpar_cookie_refresh(response), 401

        cursor.execute(
            """
            UPDATE auth_sessions
            SET last_used_at = now(), rotated_at = now()
            WHERE id = %s
            """,
            (current_session["id"],),
        )
        next_session, next_refresh_token = criar_auth_session(
            cursor,
            current_session["usuario_id"],
            current_session,
            current_session["familia_id"],
        )
        conn.commit()
    except Exception:
        conn.rollback()
        return jsonify({"erro": "Não foi possível renovar a sessão."}), 500
    finally:
        cursor.close()
        conn.close()

    access_token = gerar_access_token(
        current_session["usuario_id"],
        current_session["funcionario_id"],
        next_session["id"],
    )
    session["auth_session_id"] = str(next_session["id"])
    return resposta_com_tokens(
        {
            "access_token": access_token,
            "refresh_token": next_refresh_token,
            "token_type": "Bearer",
            "expires_in": current_app.config["JWT_ACCESS_TOKEN_MINUTES"] * 60,
        },
        next_refresh_token,
    )


@views_bp.get("/auth/me")
@access_token_required
def auth_me():
    contexto = g.auth_context

    return (
        jsonify(
            {
                "user_id": str(contexto["user_id"]),
                "funcionario_id": str(contexto["funcionario_id"]),
                "empresa_id": str(contexto["empresa_id"]),
            }
        ),
        200,
    )


@views_bp.post("/auth/logout")
@access_token_required
def auth_logout():
    try:
        revogar_familia_por_sessao(g.auth_session_id)
    except Exception:
        return jsonify({"erro": "Não foi possível encerrar a sessão."}), 500

    session.clear()
    response = limpar_cookie_refresh(jsonify({"ok": True}))
    return response, 200


# ==================================================================================================
# FUNÇÃO PRINCIPAL - CADASTRO DE USUÁRIOS
# ==================================================================================================
@views_bp.route("/cadastrar_usuario", methods=["POST"])
@require_roles("administrador")
def cadastrar_usuario():
    conn = None
    cursor = None

    try:
        dados = validar_dados_administrativos(
            request.get_json(silent=True) or {},
            password_required=True,
        )
        empresa_id_autorizada(dados["empresa_id"])

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cursor.execute(
            "SELECT cpf, email FROM usuarios WHERE cpf = %s OR email = %s",
            (dados["cpf"], dados["email"]),
        )
        duplicate = cursor.fetchone()
        if duplicate:
            message = (
                "CPF já cadastrado."
                if duplicate["cpf"] == dados["cpf"]
                else "E-mail já cadastrado."
            )
            conn.rollback()
            return jsonify({"status": "erro", "mensagem": message}), 409

        validar_estrutura_funcional(cursor, dados)

        cursor.execute(
            """
            INSERT INTO usuarios (nome, email, telefone, cpf, senha_hash)
            VALUES (%s, %s, %s, %s, %s)
            RETURNING id
            """,
            (
                dados["nome"],
                dados["email"],
                dados["telefone"],
                dados["cpf"],
                generate_password_hash(dados["senha"]),
            ),
        )

        usuario_id = cursor.fetchone()["id"]

        cursor.execute(
            """
            INSERT INTO funcionarios (
                usuario_id, empresa_id, unidade_id, equipe_id, cargo_id,
                matricula, data_admissao, tipo_contrato,
                carga_horaria_semanal, perfil, status
            )
            VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s,'ativo')
            """,
            (
                usuario_id,
                dados["empresa_id"],
                dados["unidade_id"],
                dados["equipe_id"],
                dados["cargo_id"],
                dados["matricula"],
                dados["data_admissao"],
                dados["tipo_contrato"],
                dados["carga_horaria_semanal"],
                dados["perfil"],
            ),
        )

        conn.commit()

        return jsonify({"status": "ok", "mensagem": "Usuário cadastrado com sucesso."}), 201

    except PermissionError as exc:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(exc)}), 403
    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(exc)}), 400
    except psycopg2.errors.UniqueViolation:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": "CPF, e-mail ou matrícula já cadastrado."}), 409
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": "Não foi possível cadastrar o usuário."}), 500

    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


# ==================================================================================================
# FUNÇÃO PRINCIPAL - CADASTRO DE EMPRESAS
# ==================================================================================================
@views_bp.route("/cadastrar_empresa", methods=["POST"])
@require_roles("administrador")
def cadastrar_empresa():
    conn = None
    cursor = None

    try:
        dadosEmp = request.get_json()

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cursor.execute(
            """
            INSERT INTO empresas_teste (cnpj, razao)
            VALUES (%s, %s)
            """,
            (dadosEmp["cnpj"], dadosEmp["razao"]),
        )

        conn.commit()

        return jsonify({"status": "ok", "mensagem": "Empresa cadastrada com sucesso"})

    except Exception as e:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(e)})

    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


# ==================================================================================================
# FUNÇÃO - CHAMADA DO MENU PRINCIPAL
# ==================================================================================================
@views_bp.route("/menu")
@login_required
def menu():
    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    cursor.execute("SELECT nome FROM usuarios WHERE id = %s", (session["user_id"],))
    cursor.fetchone()

    cursor.close()
    conn.close()

    return render_template(
        "menu.html",
        nome=session.get("nome"),
        tipo=session.get("tipo"),
    )


# ==================================================================================================
# FUNÇÃO - RENDERIZAÇÃO DO MENU JÁ AUTENTICADO
# ==================================================================================================
@views_bp.route("/login-page")
def login_page():
    if "user_id" in session:
        return redirect("/menu")
    return render_template("login.html")


# ==================================================================================================
# FUNÇÃO - LOGOUT
# ==================================================================================================
@views_bp.route("/logout")
def logout():
    auth_session_id = session.get("auth_session_id")
    if auth_session_id:
        try:
            revogar_familia_por_sessao(auth_session_id)
        except Exception:
            return "Não foi possível encerrar a sessão.", 500
    session.clear()
    response = limpar_cookie_refresh(redirect("/login-page"))
    return response


# ==================================================================================================
# FUNÇÃO - MEU PERFIL
# ==================================================================================================
@views_bp.route("/perfil")
@require_roles("administrador", "funcionario", "gestor", "rh")
def perfil():
    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        cursor.execute(
            """
            SELECT
                u.nome,
                u.cpf,
                u.email,
                u.telefone,
                u.status AS usuario_status,
                f.matricula,
                f.perfil,
                f.status AS funcionario_status,
                f.data_admissao,
                f.tipo_contrato,
                f.carga_horaria_semanal,
                e.nome AS empresa_nome,
                un.nome AS unidade_nome,
                eq.nome AS equipe_nome,
                c.nome AS cargo_nome
            FROM usuarios u
            INNER JOIN funcionarios f
                ON f.usuario_id = u.id
               AND f.id = %s
               AND f.empresa_id = %s
            INNER JOIN empresas e
                ON e.id = f.empresa_id
            INNER JOIN unidades un
                ON un.id = f.unidade_id
               AND un.empresa_id = f.empresa_id
            LEFT JOIN equipes eq
                ON eq.id = f.equipe_id
               AND eq.empresa_id = f.empresa_id
               AND eq.unidade_id = f.unidade_id
            LEFT JOIN cargos c
                ON c.id = f.cargo_id
               AND c.empresa_id = f.empresa_id
            WHERE u.id = %s
            """,
            (
                g.auth_funcionario_id,
                g.auth_context["empresa_id"],
                g.auth_user_id,
            ),
        )
        user = cursor.fetchone()
    finally:
        cursor.close()
        conn.close()

    if not user:
        return "Perfil não encontrado", 404

    return render_template("perfil.html", user=user)


# ==================================================================================================
# FUNÇÃO - LISTAGEM DE USUÁRIOS (GERENCIAMENTO)
# ==================================================================================================
@views_bp.route("/listarUsuarios")
@require_roles("administrador")
def listar_usuarios():
    try:
        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cursor.execute(
            """
            SELECT
                u.id,
                f.id AS funcionario_id,
                u.nome,
                u.cpf,
                u.email,
                u.telefone,
                f.status,
                f.perfil,
                f.matricula,
                e.nome AS empresa,
                un.nome AS unidade,
                eq.nome AS equipe,
                c.nome AS cargo
            FROM usuarios u
            INNER JOIN funcionarios f ON f.usuario_id = u.id
            INNER JOIN empresas e ON e.id = f.empresa_id
            INNER JOIN unidades un ON un.id = f.unidade_id
            LEFT JOIN equipes eq ON eq.id = f.equipe_id
            LEFT JOIN cargos c ON c.id = f.cargo_id
            WHERE f.empresa_id = %s
            ORDER BY u.nome, e.nome
            """
            ,
            (g.auth_context["empresa_id"],),
        )

        usuarios = cursor.fetchall()

        cursor.close()
        conn.close()

        return jsonify(usuarios)

    except Exception:
        return jsonify([]), 500


# ==================================================================================================
# FUNÇÃO - LISTAGEM DE EMPRESAS (GERENCIAMENTO)
# ==================================================================================================
@views_bp.route("/listarEmpresas")
@require_roles("administrador")
def listar_empresas():
    try:
        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cursor.execute(
            """
            SELECT id, nome
            FROM empresas
            WHERE id = %s AND status = 'ativa'
            ORDER BY nome
            """,
            (g.auth_context["empresa_id"],),
        )
        empresas = cursor.fetchall()

        cursor.close()
        conn.close()

        return jsonify(empresas)

    except Exception:
        return jsonify([]), 500


# ==================================================================================================
# FUNÇÃO - ALTERAÇÃO DE DADOS DO USUÁRIO
# ==================================================================================================
@views_bp.route("/atualizar_usuario", methods=["POST"])
@require_roles("administrador")
def atualizar_usuario():
    conn = None
    cursor = None
    try:
        request_data = request.get_json(silent=True) or {}
        funcionario_id = uuid_obrigatorio(
            request_data.get("funcionario_id"),
            "Funcionário",
        )
        dados = validar_dados_administrativos(request_data)
        empresa_id_autorizada(dados["empresa_id"])

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        cursor.execute(
            "SELECT usuario_id FROM funcionarios WHERE id = %s AND empresa_id = %s FOR UPDATE",
            (funcionario_id, g.auth_context["empresa_id"]),
        )
        funcionario = cursor.fetchone()
        if not funcionario:
            conn.rollback()
            return jsonify({"status": "erro", "mensagem": "Funcionário não encontrado."}), 404

        user_id = funcionario["usuario_id"]
        cursor.execute(
            """
            SELECT cpf, email
            FROM usuarios
            WHERE id <> %s AND (cpf = %s OR email = %s)
            """,
            (user_id, dados["cpf"], dados["email"]),
        )
        duplicate = cursor.fetchone()
        if duplicate:
            message = (
                "CPF já cadastrado."
                if duplicate["cpf"] == dados["cpf"]
                else "E-mail já cadastrado."
            )
            conn.rollback()
            return jsonify({"status": "erro", "mensagem": message}), 409

        validar_estrutura_funcional(cursor, dados)

        cursor.execute(
            """
            UPDATE usuarios
            SET nome = %s, email = %s, telefone = %s, cpf = %s, updated_at = now()
            WHERE id = %s
            """,
            (dados["nome"], dados["email"], dados["telefone"], dados["cpf"], user_id),
        )

        cursor.execute(
            """
            UPDATE funcionarios
            SET empresa_id = %s,
                unidade_id = %s,
                equipe_id = %s,
                cargo_id = %s,
                matricula = %s,
                tipo_contrato = %s,
                data_admissao = %s,
                carga_horaria_semanal = %s,
                perfil = %s,
                updated_at = now()
            WHERE id = %s
            """,
            (
                dados["empresa_id"],
                dados["unidade_id"],
                dados["equipe_id"],
                dados["cargo_id"],
                dados["matricula"],
                dados["tipo_contrato"],
                dados["data_admissao"],
                dados["carga_horaria_semanal"],
                dados["perfil"],
                funcionario_id,
            ),
        )

        conn.commit()

        return jsonify({"status": "ok"})

    except PermissionError as exc:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(exc)}), 403
    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(exc)}), 400
    except psycopg2.errors.UniqueViolation:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": "CPF, e-mail ou matrícula já cadastrado."}), 409
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": "Não foi possível atualizar o usuário."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


# ==================================================================================================
# FUNÇÃO - ALTERAÇÃO DE HORÁRIOS DO USUÁRIO
# ==================================================================================================
@views_bp.route("/atualizar_horarios", methods=["POST"])
@require_roles("administrador")
def atualizar_horarios():
    return (
        jsonify(
            {
                "status": "erro",
                "mensagem": (
                    "A interface antiga de horários não suporta múltiplos períodos. "
                    "Utilize a API de turnos."
                ),
            }
        ),
        410,
    )


@views_bp.route("/api/admin/turnos", methods=["GET", "POST"])
@require_roles("administrador")
def turnos_administrativos():
    conn = None
    cursor = None
    try:
        empresa_id = empresa_autenticada(
            request.get_json(silent=True) if request.method == "POST" else None
        )
        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

        if request.method == "GET":
            if request.args.get("empresa_id") not in (None, ""):
                empresa_id_autorizada(request.args["empresa_id"])
            cursor.execute(
                """
                SELECT id
                FROM turnos
                WHERE empresa_id = %s
                ORDER BY nome, id
                """,
                (empresa_id,),
            )
            turnos = []
            for item in cursor.fetchall():
                turno = buscar_turno(cursor, empresa_id, item["id"])
                turnos.append(serializar_turno(turno))
            return jsonify({"turnos": turnos}), 200

        dados = request.get_json(silent=True) or {}
        turno_data = validar_turno_payload(dados)
        cursor.execute(
            """
            INSERT INTO turnos (empresa_id, nome, timezone, status)
            VALUES (%s, %s, %s, %s)
            RETURNING id
            """,
            (
                empresa_id,
                turno_data["nome"],
                turno_data["timezone"],
                turno_data["status"],
            ),
        )
        turno_id = cursor.fetchone()["id"]
        inserir_periodos_turno(
            cursor, empresa_id, turno_id, turno_data["periodos"]
        )
        turno = buscar_turno(cursor, empresa_id, turno_id)
        conn.commit()
        return jsonify({"turno": serializar_turno(turno)}), 201
    except PermissionError as exc:
        if conn:
            conn.rollback()
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"erro": str(exc)}), 400
    except psycopg2.errors.UniqueViolation:
        if conn:
            conn.rollback()
        return jsonify({"erro": "Já existe um turno com esses dados."}), 409
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"erro": "Não foi possível processar o turno."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@views_bp.route("/api/admin/turnos/<turno_id>", methods=["GET", "PUT"])
@require_roles("administrador")
def turno_administrativo(turno_id):
    conn = None
    cursor = None
    try:
        turno_id = uuid_obrigatorio(turno_id, "Turno")
        dados = request.get_json(silent=True) or {}
        empresa_id = empresa_autenticada(dados if request.method == "PUT" else None)
        if request.args.get("empresa_id") not in (None, ""):
            empresa_id_autorizada(request.args["empresa_id"])

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        turno = buscar_turno(
            cursor, empresa_id, turno_id, bloquear=request.method == "PUT"
        )
        if not turno:
            return jsonify({"erro": "Turno não encontrado."}), 404
        if request.method == "GET":
            return jsonify({"turno": serializar_turno(turno)}), 200

        turno_data = validar_turno_payload(dados, partial=True)
        if "periodos" in turno_data:
            cursor.execute(
                """
                SELECT 1
                FROM funcionarios_turnos
                WHERE empresa_id = %s AND turno_id = %s
                LIMIT 1
                """,
                (empresa_id, turno_id),
            )
            if cursor.fetchone():
                return (
                    jsonify(
                        {
                            "erro": (
                                "Períodos de um turno já atribuído não podem ser "
                                "alterados; crie um novo turno para preservar o histórico."
                            )
                        }
                    ),
                    409,
                )

        cursor.execute(
            """
            UPDATE turnos
            SET nome = %s, timezone = %s, status = %s, updated_at = now()
            WHERE id = %s AND empresa_id = %s
            """,
            (
                turno_data.get("nome", turno["nome"]),
                turno_data.get("timezone", turno["timezone"]),
                turno_data.get("status", turno["status"]),
                turno_id,
                empresa_id,
            ),
        )
        if "periodos" in turno_data:
            cursor.execute(
                "DELETE FROM periodos_turno WHERE turno_id = %s AND empresa_id = %s",
                (turno_id, empresa_id),
            )
            inserir_periodos_turno(
                cursor, empresa_id, turno_id, turno_data["periodos"]
            )
        updated = buscar_turno(cursor, empresa_id, turno_id)
        conn.commit()
        return jsonify({"turno": serializar_turno(updated)}), 200
    except PermissionError as exc:
        if conn:
            conn.rollback()
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"erro": str(exc)}), 400
    except psycopg2.errors.UniqueViolation:
        if conn:
            conn.rollback()
        return jsonify({"erro": "Já existe um turno com esses dados."}), 409
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"erro": "Não foi possível atualizar o turno."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


# ==================================================================================================
# FUNÇÕES - ENVIO E VALIDAÇÃO DE TOKEN PARA RECUPERAÇÃO DE SENHA
# ==================================================================================================
PASSWORD_RESET_PUBLIC_MESSAGE = (
    "Se os dados corresponderem a uma conta elegível, as instruções de redefinição serão disponibilizadas."
)
PASSWORD_RESET_INVALID_MESSAGE = "Token inválido ou expirado."


def _password_reset_user(cursor, identifier):
    identifier = str(identifier or "").strip()
    if "@" in identifier:
        cursor.execute(
            """
            SELECT id, senha_hash, status
            FROM usuarios
            WHERE email = %s AND status IN ('ativo', 'bloqueado')
            """,
            (identifier.lower(),),
        )
    else:
        cpf = normalizar_cpf(identifier)
        if len(cpf) != 11:
            return None
        cursor.execute(
            """
            SELECT id, senha_hash, status
            FROM usuarios
            WHERE cpf = %s AND status IN ('ativo', 'bloqueado')
            """,
            (cpf,),
        )
    return cursor.fetchone()


@views_bp.route("/auth/password/forgot", methods=["POST"])
@views_bp.route("/enviar-token", methods=["POST"])
def enviar_token():
    data = request.get_json(silent=True) or {}
    identifier = data.get("identifier") or data.get("email") or data.get("cpf")
    raw_token = secrets.token_urlsafe(32)
    token_hash = hashlib.sha256(raw_token.encode("utf-8")).hexdigest()
    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        user = _password_reset_user(cursor, identifier)
        if user:
            purpose = "first_access" if not user["senha_hash"] else "password_reset"
            cursor.execute(
                """
                UPDATE password_reset_tokens
                SET revoked_at = now()
                WHERE usuario_id = %s
                  AND used_at IS NULL
                  AND revoked_at IS NULL
                """,
                (user["id"],),
            )
            cursor.execute(
                """
                INSERT INTO password_reset_tokens (
                    usuario_id, token_hash, purpose, expires_at
                )
                VALUES (
                    %s, %s, %s,
                    now() + (%s * interval '1 minute')
                )
                """,
                (
                    user["id"],
                    token_hash,
                    purpose,
                    current_app.config["PASSWORD_RESET_TOKEN_MINUTES"],
                ),
            )
        conn.commit()
    except Exception:
        conn.rollback()
        return jsonify({"erro": "Não foi possível processar a solicitação."}), 500
    finally:
        cursor.close()
        conn.close()

    response = {"ok": True, "mensagem": PASSWORD_RESET_PUBLIC_MESSAGE}
    configured_test_key = current_app.config.get("PASSWORD_RESET_TEST_KEY", "")
    supplied_test_key = request.headers.get("X-Password-Reset-Test-Key", "")
    if (
        user
        and configured_test_key
        and supplied_test_key
        and secrets.compare_digest(configured_test_key, supplied_test_key)
    ):
        response["test_token"] = raw_token
    return jsonify(response), 200


@views_bp.route("/auth/password/validate", methods=["POST"])
@views_bp.route("/validar-token", methods=["POST"])
def validar_token():
    data = request.get_json(silent=True) or {}
    token = data.get("token")
    if not isinstance(token, str) or not token:
        return jsonify({"erro": PASSWORD_RESET_INVALID_MESSAGE}), 400
    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    conn = conectar_bd()
    cursor = conn.cursor()
    try:
        cursor.execute(
            """
            SELECT 1
            FROM password_reset_tokens
            WHERE token_hash = %s
              AND used_at IS NULL
              AND revoked_at IS NULL
              AND expires_at > now()
            """,
            (token_hash,),
        )
        valid = cursor.fetchone() is not None
    finally:
        cursor.close()
        conn.close()
    if not valid:
        return jsonify({"erro": PASSWORD_RESET_INVALID_MESSAGE}), 400
    return jsonify({"ok": True}), 200


# ==================================================================================================
# FUNÇÃO - REDEFINIÇÃO DE SENHA
# ==================================================================================================
@views_bp.route("/auth/password/reset", methods=["POST"])
@views_bp.route("/resetar-senha", methods=["POST"])
def resetar_senha():
    data = request.get_json(silent=True) or {}
    token = data.get("token")
    senha = data.get("new_password") or data.get("senha")
    if not isinstance(senha, str) or len(senha) < 12:
        return jsonify({"erro": "A senha deve possuir pelo menos 12 caracteres."}), 400
    if not isinstance(token, str) or not token:
        return jsonify({"erro": PASSWORD_RESET_INVALID_MESSAGE}), 400

    token_hash = hashlib.sha256(token.encode("utf-8")).hexdigest()
    senha_hash = generate_password_hash(senha)
    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    try:
        cursor.execute(
            """
            SELECT id, usuario_id, purpose
            FROM password_reset_tokens
            WHERE token_hash = %s
              AND used_at IS NULL
              AND revoked_at IS NULL
              AND expires_at > now()
            FOR UPDATE
            """,
            (token_hash,),
        )
        reset_token = cursor.fetchone()
        if not reset_token:
            conn.rollback()
            return jsonify({"erro": PASSWORD_RESET_INVALID_MESSAGE}), 400

        cursor.execute(
            """
            UPDATE usuarios
            SET senha_hash = %s,
                status = CASE
                    WHEN %s = 'first_access' AND status = 'bloqueado' THEN 'ativo'
                    ELSE status
                END,
                updated_at = now()
            WHERE id = %s
            """,
            (senha_hash, reset_token["purpose"], reset_token["usuario_id"]),
        )
        if reset_token["purpose"] == "first_access":
            cursor.execute(
                """
                UPDATE funcionarios
                SET status = 'ativo', updated_at = now()
                WHERE usuario_id = %s AND status = 'afastado'
                """,
                (reset_token["usuario_id"],),
            )
        cursor.execute(
            "UPDATE password_reset_tokens SET used_at = now() WHERE id = %s",
            (reset_token["id"],),
        )
        cursor.execute(
            """
            UPDATE password_reset_tokens
            SET revoked_at = now()
            WHERE usuario_id = %s
              AND id <> %s
              AND used_at IS NULL
              AND revoked_at IS NULL
            """,
            (reset_token["usuario_id"], reset_token["id"]),
        )
        conn.commit()
    except Exception:
        conn.rollback()
        return jsonify({"erro": "Não foi possível redefinir a senha."}), 500
    finally:
        cursor.close()
        conn.close()
    session.clear()
    return jsonify({"ok": True}), 200


# ==================================================================================================
# FUNÇÃO - LISTAGEM DE USUÁRIOS NO CADASTRO DE FOTOS (APENAS USUÁRIOS SEM FOTO)
# ==================================================================================================
@views_bp.route("/listar_usuarios_select", methods=["GET"])
def listar_usuarios_select():
    try:
        conn = conectar_bd()
        cursor = conn.cursor()

        cursor.execute("""
            SELECT u.id, u.nome
            FROM usuarios u
            WHERE NOT EXISTS (
                SELECT 1 FROM fotos f WHERE f.nome = u.nome
            )
            """)

        usuarios = cursor.fetchall()

        lista = [{"id": u[0], "nome": u[1]} for u in usuarios]

        cursor.close()
        conn.close()

        return jsonify(lista)

    except Exception as e:
        print("ERRO:", e)
        return jsonify([])


# ==================================================================================================
# FUNÇÃO - RETORNO DO HORÁRIO DO SERVIDOR
# ==================================================================================================
@views_bp.route("/hora-servidor")
def hora_servidor():
    agora = datetime.now(ZoneInfo("America/Sao_Paulo"))

    dias = ["Seg", "Ter", "Qua", "Qui", "Sex", "Sáb", "Dom"]
    dia_semana = dias[agora.weekday()]

    return {
        "data": agora.strftime("%d/%m/%Y"),
        "hora": agora.strftime("%H:%M:%S"),
        "dia": dia_semana,
    }


# ==================================================================================================
# JORNADAS E ATRIBUIÇÕES DE TURNOS
# ==================================================================================================
@views_bp.route(
    "/api/admin/funcionarios/<funcionario_id>/turnos", methods=["GET", "POST"]
)
@require_roles("administrador")
def turnos_funcionario_administrativo(funcionario_id):
    conn = None
    cursor = None
    try:
        funcionario_id = uuid_obrigatorio(funcionario_id, "Funcionário")
        dados = request.get_json(silent=True) or {}
        empresa_id = empresa_autenticada(dados if request.method == "POST" else None)
        if request.args.get("empresa_id") not in (None, ""):
            empresa_id_autorizada(request.args["empresa_id"])
        if dados.get("funcionario_id") not in (None, ""):
            body_funcionario_id = uuid_obrigatorio(
                dados["funcionario_id"], "Funcionário"
            )
            if body_funcionario_id != funcionario_id:
                raise PermissionError("Funcionário fora do escopo solicitado.")

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute(
            f"""
            SELECT id, status
            FROM funcionarios
            WHERE id = %s AND empresa_id = %s
            {"FOR UPDATE" if request.method == "POST" else ""}
            """,
            (funcionario_id, empresa_id),
        )
        funcionario = cursor.fetchone()
        if not funcionario:
            return jsonify({"erro": "Funcionário não encontrado."}), 404

        if request.method == "GET":
            cursor.execute(
                """
                SELECT
                    ft.id, ft.vigencia_inicio, ft.vigencia_fim,
                    t.id AS turno_id, t.nome AS turno_nome
                FROM funcionarios_turnos ft
                INNER JOIN turnos t
                    ON t.id = ft.turno_id
                   AND t.empresa_id = ft.empresa_id
                WHERE ft.funcionario_id = %s AND ft.empresa_id = %s
                ORDER BY ft.vigencia_inicio, ft.created_at, ft.id
                """,
                (funcionario_id, empresa_id),
            )
            historico = [
                {
                    "id": str(item["id"]),
                    "turno": {
                        "id": str(item["turno_id"]),
                        "nome": item["turno_nome"],
                    },
                    "vigencia": {
                        "inicio": item["vigencia_inicio"].isoformat(),
                        "fim": (
                            item["vigencia_fim"].isoformat()
                            if item["vigencia_fim"]
                            else None
                        ),
                    },
                }
                for item in cursor.fetchall()
            ]
            return jsonify({"historico": historico}), 200

        if funcionario["status"] != "ativo":
            return jsonify({"erro": "O vínculo do funcionário não está ativo."}), 409

        turno_id = uuid_obrigatorio(dados.get("turno_id"), "Turno")
        vigencia_inicio = parse_iso_date(
            dados.get("vigencia_inicio"), "Início da vigência"
        )
        vigencia_fim = parse_optional_iso_date(
            dados.get("vigencia_fim"), "Fim da vigência"
        )
        if vigencia_fim and vigencia_fim < vigencia_inicio:
            raise ValueError("O fim da vigência não pode anteceder o início.")

        cursor.execute(
            """
            SELECT id
            FROM turnos
            WHERE id = %s AND empresa_id = %s AND status = 'ativo'
            """,
            (turno_id, empresa_id),
        )
        if not cursor.fetchone():
            return jsonify({"erro": "Turno ativo não encontrado."}), 404

        cursor.execute(
            """
            SELECT id, vigencia_inicio, vigencia_fim
            FROM funcionarios_turnos
            WHERE funcionario_id = %s AND empresa_id = %s
            ORDER BY vigencia_inicio, id
            FOR UPDATE
            """,
            (funcionario_id, empresa_id),
        )
        atribuicoes = cursor.fetchall()
        anterior = None
        for atribuicao in atribuicoes:
            inicio_existente = atribuicao["vigencia_inicio"]
            fim_existente = atribuicao["vigencia_fim"]
            if inicio_existente == vigencia_inicio:
                return (
                    jsonify({"erro": "Já existe atribuição iniciada nessa data."}),
                    409,
                )
            if inicio_existente > vigencia_inicio and (
                vigencia_fim is None or inicio_existente <= vigencia_fim
            ):
                return (
                    jsonify({"erro": "A vigência informada sobrepõe atribuição futura."}),
                    409,
                )
            if inicio_existente < vigencia_inicio and (
                fim_existente is None or fim_existente >= vigencia_inicio
            ):
                if anterior:
                    return jsonify({"erro": "O histórico atual possui sobreposição."}), 409
                anterior = atribuicao

        if anterior:
            cursor.execute(
                """
                UPDATE funcionarios_turnos
                SET vigencia_fim = %s
                WHERE id = %s AND empresa_id = %s
                """,
                (vigencia_inicio - timedelta(days=1), anterior["id"], empresa_id),
            )

        cursor.execute(
            """
            INSERT INTO funcionarios_turnos (
                empresa_id, funcionario_id, turno_id,
                vigencia_inicio, vigencia_fim
            )
            VALUES (%s, %s, %s, %s, %s)
            RETURNING id
            """,
            (
                empresa_id,
                funcionario_id,
                turno_id,
                vigencia_inicio,
                vigencia_fim,
            ),
        )
        atribuicao_id = cursor.fetchone()["id"]
        conn.commit()
        return (
            jsonify(
                {
                    "atribuicao": {
                        "id": str(atribuicao_id),
                        "funcionario_id": funcionario_id,
                        "turno_id": turno_id,
                        "vigencia": {
                            "inicio": vigencia_inicio.isoformat(),
                            "fim": vigencia_fim.isoformat() if vigencia_fim else None,
                        },
                    }
                }
            ),
            201,
        )
    except PermissionError as exc:
        if conn:
            conn.rollback()
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"erro": str(exc)}), 400
    except psycopg2.errors.UniqueViolation:
        if conn:
            conn.rollback()
        return jsonify({"erro": "Já existe atribuição iniciada nessa data."}), 409
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"erro": "Não foi possível processar a atribuição."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@views_bp.route(
    "/api/admin/funcionarios/<funcionario_id>/jornada", methods=["GET"]
)
@require_roles("administrador")
def jornada_funcionario_administrativo(funcionario_id):
    conn = None
    cursor = None
    try:
        funcionario_id = uuid_obrigatorio(funcionario_id, "Funcionário")
        empresa_id = empresa_autenticada()
        if request.args.get("empresa_id") not in (None, ""):
            empresa_id_autorizada(request.args["empresa_id"])
        data_consulta = (
            parse_iso_date(request.args["data"], "Data")
            if request.args.get("data")
            else date.today()
        )
        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute(
            "SELECT 1 FROM funcionarios WHERE id = %s AND empresa_id = %s",
            (funcionario_id, empresa_id),
        )
        if not cursor.fetchone():
            return jsonify({"erro": "Funcionário não encontrado."}), 404
        jornada = buscar_jornada_data(
            cursor, empresa_id, funcionario_id, data_consulta
        )
        return jsonify({"data": data_consulta.isoformat(), "jornada": jornada}), 200
    except PermissionError as exc:
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400
    except Exception:
        return jsonify({"erro": "Não foi possível consultar a jornada."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@views_bp.route("/api/jornada", methods=["GET"])
@views_bp.route("/jornada", methods=["GET"])
@require_roles("administrador", "funcionario", "gestor", "rh")
def get_jornada():
    conn = None
    cursor = None
    try:
        if request.args.get("funcionario_id") or request.args.get("user_id"):
            return jsonify({"erro": "A jornada pessoal não aceita outro vínculo."}), 400
        if request.args.get("empresa_id") not in (None, ""):
            empresa_id_autorizada(request.args["empresa_id"])
        data_consulta = (
            parse_iso_date(request.args["data"], "Data")
            if request.args.get("data")
            else date.today()
        )
        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        jornada = buscar_jornada_data(
            cursor,
            g.auth_context["empresa_id"],
            g.auth_funcionario_id,
            data_consulta,
        )
        return jsonify({"data": data_consulta.isoformat(), "jornada": jornada}), 200
    except PermissionError as exc:
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400
    except Exception:
        return jsonify({"erro": "Não foi possível consultar a jornada."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


def formatar_hora(hora):
    return hora.strftime("%H:%M") if hora else "--:--"


def calcular_total(entrada, saida, saida_intervalo=None, volta_intervalo=None):
    try:
        if not entrada or not saida:
            return "--"

        dt_entrada = datetime.combine(date.today(), entrada)
        dt_saida = datetime.combine(date.today(), saida)

        total = dt_saida - dt_entrada

        if saida_intervalo and volta_intervalo:
            dt_saida_intervalo = datetime.combine(date.today(), saida_intervalo)
            dt_volta_intervalo = datetime.combine(date.today(), volta_intervalo)
            total -= dt_volta_intervalo - dt_saida_intervalo

        total_segundos = int(total.total_seconds())

        if total_segundos < 0:
            return "--"

        horas_total = total_segundos // 3600
        minutos_total = (total_segundos % 3600) // 60

        return f"{horas_total}h{minutos_total:02d}"
    except Exception:
        return "--"


def montar_registros_ponto(
    user_id, data_inicio=None, data_fim=None, formato_data="iso"
):
    conn = None
    cursor = None

    try:
        conn = conectar_bd()
        cursor = conn.cursor()

        query = """
            SELECT data_registro, horario_registro
            FROM ponto
            WHERE usuario_id = %s
        """
        params = [user_id]

        if data_inicio:
            query += " AND data_registro >= %s"
            params.append(data_inicio)

        if data_fim:
            query += " AND data_registro <= %s"
            params.append(data_fim)

        query += " ORDER BY data_registro ASC, horario_registro ASC"

        cursor.execute(query, tuple(params))
        registros = cursor.fetchall()

        dias = defaultdict(list)

        for data_registro, horario_registro in registros:
            dias[data_registro].append(horario_registro)

        mapa_dias = {
            "Monday": "Segunda",
            "Tuesday": "Terça",
            "Wednesday": "Quarta",
            "Thursday": "Quinta",
            "Friday": "Sexta",
            "Saturday": "Sábado",
            "Sunday": "Domingo",
        }

        resultado = []

        for data_registro in sorted(dias.keys()):
            horarios = dias[data_registro]

            entrada = None
            saida_intervalo = None
            volta_intervalo = None
            saida = None

            if len(horarios) == 1:
                entrada = horarios[0]
            elif len(horarios) == 2:
                entrada = horarios[0]
                saida = horarios[1]
            elif len(horarios) == 3:
                entrada = horarios[0]
                saida_intervalo = horarios[1]
                saida = horarios[2]
            elif len(horarios) >= 4:
                entrada = horarios[0]
                saida_intervalo = horarios[1]
                volta_intervalo = horarios[2]
                saida = horarios[3]

            nome_dia_en = data_registro.strftime("%A")

            if formato_data == "br":
                data_formatada = data_registro.strftime("%d/%m/%Y")
            else:
                data_formatada = data_registro.strftime("%Y-%m-%d")

            resultado.append(
                {
                    "data": data_formatada,
                    "dia": mapa_dias.get(nome_dia_en, nome_dia_en),
                    "entrada": formatar_hora(entrada),
                    "saida_intervalo": formatar_hora(saida_intervalo),
                    "volta_intervalo": formatar_hora(volta_intervalo),
                    "saida": formatar_hora(saida),
                    "total": calcular_total(
                        entrada,
                        saida,
                        saida_intervalo,
                        volta_intervalo,
                    ),
                }
            )

        return resultado

    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()


@views_bp.route("/pontos", methods=["GET"])
@login_required
def listar_pontos():
    try:
        usuario_id = request.args.get("usuario_id")
        data_inicio = request.args.get("inicio")
        data_fim = request.args.get("fim")

        # se for administrador e escolheu alguém, usa o usuário escolhido
        if session.get("tipo") == "admin" and usuario_id:
            user_id = usuario_id
        else:
            user_id = session["user_id"]

        resultado = montar_registros_ponto(
            user_id=user_id,
            data_inicio=data_inicio,
            data_fim=data_fim,
            formato_data="iso",
        )

        return jsonify(resultado), 200

    except Exception as e:
        return jsonify({"erro": str(e)}), 500


@views_bp.route("/exportar-pontos", methods=["GET"])
@login_required
def exportar_pontos():
    user_id = session["user_id"]

    formato = request.args.get("formato", "").lower()
    data_inicio = request.args.get("inicio")
    data_fim = request.args.get("fim")

    registros = montar_registros_ponto(
        user_id=user_id,
        data_inicio=data_inicio,
        data_fim=data_fim,
        formato_data="br",
    )

    if not registros:
        return jsonify({"erro": "Nenhum registro encontrado para exportação."}), 404

    if formato == "csv":
        output = io.StringIO()
        writer = csv.writer(output, delimiter=";")
        writer.writerow(
            ["Dia", "Data", "Entrada", "Saída Int.", "Volta Int.", "Saída", "Total"]
        )

        for item in registros:
            writer.writerow(
                [
                    item["dia"],
                    item["data"],
                    item["entrada"],
                    item["saida_intervalo"],
                    item["volta_intervalo"],
                    item["saida"],
                    item["total"],
                ]
            )

        mem = io.BytesIO()
        mem.write(output.getvalue().encode("utf-8-sig"))
        mem.seek(0)

        return send_file(
            mem,
            as_attachment=True,
            download_name="controle_ponto.csv",
            mimetype="text/csv",
        )

    elif formato == "pdf":
        mem = io.BytesIO()
        pdf = canvas.Canvas(mem, pagesize=A4)
        largura, altura = A4

        y = altura - 40
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(40, y, "Relatório de Controle de Ponto")

        y -= 30
        pdf.setFont("Helvetica", 9)
        pdf.drawString(40, y, "Dia")
        pdf.drawString(95, y, "Data")
        pdf.drawString(160, y, "Entrada")
        pdf.drawString(220, y, "Saída Int.")
        pdf.drawString(295, y, "Volta Int.")
        pdf.drawString(370, y, "Saída")
        pdf.drawString(430, y, "Total")

        y -= 15
        pdf.line(40, y, 550, y)

        for item in registros:
            y -= 18

            if y < 50:
                pdf.showPage()
                y = altura - 40
                pdf.setFont("Helvetica", 9)

            pdf.drawString(40, y, str(item["dia"]))
            pdf.drawString(95, y, str(item["data"]))
            pdf.drawString(160, y, str(item["entrada"]))
            pdf.drawString(220, y, str(item["saida_intervalo"]))
            pdf.drawString(295, y, str(item["volta_intervalo"]))
            pdf.drawString(370, y, str(item["saida"]))
            pdf.drawString(430, y, str(item["total"]))

        pdf.save()
        mem.seek(0)

        return send_file(
            mem,
            as_attachment=True,
            download_name="controle_ponto.pdf",
            mimetype="application/pdf",
        )

    elif formato in ["word", "doc"]:
        doc = Document()
        doc.add_heading("Relatório de Controle de Ponto", level=1)

        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"

        hdr = table.rows[0].cells
        hdr[0].text = "Dia"
        hdr[1].text = "Data"
        hdr[2].text = "Entrada"
        hdr[3].text = "Saída Int."
        hdr[4].text = "Volta Int."
        hdr[5].text = "Saída"
        hdr[6].text = "Total"

        for item in registros:
            row = table.add_row().cells
            row[0].text = item["dia"]
            row[1].text = item["data"]
            row[2].text = item["entrada"]
            row[3].text = item["saida_intervalo"]
            row[4].text = item["volta_intervalo"]
            row[5].text = item["saida"]
            row[6].text = item["total"]

        mem = io.BytesIO()
        doc.save(mem)
        mem.seek(0)

        return send_file(
            mem,
            as_attachment=True,
            download_name="controle_ponto.docx",
            mimetype="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

    return jsonify({"erro": "Formato inválido."}), 400


# ==================================================================================================
# FUNÇÃO - LISTAGEM DE SETORES DO BANCO DE DADOS
# ==================================================================================================
@views_bp.route("/listar_unidades", methods=["GET"])
@require_roles("administrador")
def listar_unidades():
    try:
        empresa_id = empresa_id_autorizada(request.args.get("empresa_id"))
    except PermissionError as exc:
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400

    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute(
        """
        SELECT id, nome
        FROM unidades
        WHERE empresa_id = %s AND status = 'ativa'
        ORDER BY nome
        """,
        (empresa_id,),
    )
    unidades = cursor.fetchall()
    cursor.close()
    conn.close()
    return jsonify(unidades)


@views_bp.route("/listar_equipes", methods=["GET"])
@views_bp.route("/listar_setores", methods=["GET"])
@require_roles("administrador")
def listar_equipes():
    try:
        empresa_id = empresa_id_autorizada(request.args.get("empresa_id"))
        unidade_id = uuid_obrigatorio(request.args.get("unidade_id"), "Unidade")
    except PermissionError as exc:
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400

    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
    cursor.execute(
        """
        SELECT id, nome
        FROM equipes
        WHERE empresa_id = %s AND unidade_id = %s AND status = 'ativa'
        ORDER BY nome
        """,
        (empresa_id, unidade_id),
    )
    equipes = cursor.fetchall()
    cursor.close()
    conn.close()
    return jsonify(equipes)


# ==================================================================================================
# FUNÇÃO - LISTAGEM DE CARGOS DO BANCO DE DADOS
# ==================================================================================================
@views_bp.route("/listar_cargos", methods=["GET"])
@require_roles("administrador")
def listar_cargos():
    try:
        empresa_id = empresa_id_autorizada(request.args.get("empresa_id"))
    except PermissionError as exc:
        return jsonify({"erro": str(exc)}), 403
    except ValueError as exc:
        return jsonify({"erro": str(exc)}), 400

    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    cursor.execute(
        """
        SELECT id, nome
        FROM cargos
        WHERE empresa_id = %s AND status = 'ativo'
        ORDER BY nome
        """,
        (empresa_id,),
    )
    cargos = cursor.fetchall()

    cursor.close()
    conn.close()

    return jsonify(cargos)


@views_bp.route("/editarUsuario")
@require_roles("administrador")
def editar_usuario():
    try:
        funcionario_id = uuid_obrigatorio(
            request.args.get("funcionario_id"),
            "Funcionário",
        )
    except ValueError as exc:
        return str(exc), 400

    conn = conectar_bd()
    cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)

    cursor.execute(
        """
        SELECT 
            u.id AS usuario_id,
            f.id AS funcionario_id,
            u.nome,
            u.email,
            u.telefone,
            u.cpf,
            u.status AS usuario_status,
            f.empresa_id,
            f.unidade_id,
            f.equipe_id,
            f.cargo_id,
            f.status,
            f.matricula,
            f.perfil,
            f.tipo_contrato,
            f.data_admissao,
            f.carga_horaria_semanal,
            e.nome AS empresa,
            un.nome AS unidade,
            c.nome AS cargo,
            eq.nome AS equipe
        FROM usuarios u
        INNER JOIN funcionarios f ON u.id = f.usuario_id
        INNER JOIN empresas e ON e.id = f.empresa_id
        INNER JOIN unidades un ON un.id = f.unidade_id
        LEFT JOIN cargos c ON c.id = f.cargo_id
        LEFT JOIN equipes eq ON eq.id = f.equipe_id
        WHERE f.id = %s AND f.empresa_id = %s
        """,
        (funcionario_id, g.auth_context["empresa_id"]),
    )

    usuario = cursor.fetchone()

    cursor.close()
    conn.close()

    if not usuario:
        return "Usuário não encontrado", 404

    return render_template("editarUsuario.html", usuario=usuario)


@views_bp.route("/editarHorarios")
@require_roles("administrador")
def editar_horarios():
    return (
        jsonify(
            {
                "erro": (
                    "A interface antiga de horários não suporta múltiplos períodos. "
                    "A gestão de jornadas está disponível pela API de turnos."
                )
            }
        ),
        410,
    )


@views_bp.route("/deletar_usuario", methods=["POST"])
@require_roles("administrador")
def deletar_usuario():
    conn = None
    cursor = None
    try:
        dados = request.get_json(silent=True) or {}
        funcionario_id = uuid_obrigatorio(dados.get("funcionario_id"), "Funcionário")

        if str(funcionario_id) == str(session.get("funcionario_id")):
            return (
                jsonify(
                    {
                        "status": "erro",
                        "mensagem": "Você não pode desativar seu próprio vínculo.",
                    }
                ),
                403,
            )

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute(
            "SELECT usuario_id FROM funcionarios WHERE id = %s AND empresa_id = %s FOR UPDATE",
            (funcionario_id, g.auth_context["empresa_id"]),
        )
        funcionario = cursor.fetchone()
        if not funcionario:
            conn.rollback()
            return jsonify({"status": "erro", "mensagem": "Funcionário não encontrado."}), 404

        cursor.execute(
            "UPDATE funcionarios SET status = 'desligado', updated_at = now() WHERE id = %s AND empresa_id = %s",
            (funcionario_id, g.auth_context["empresa_id"]),
        )
        cursor.execute(
            "SELECT 1 FROM funcionarios WHERE usuario_id = %s AND status = 'ativo' LIMIT 1",
            (funcionario["usuario_id"],),
        )
        if not cursor.fetchone():
            cursor.execute(
                "UPDATE usuarios SET status = 'inativo', updated_at = now() WHERE id = %s",
                (funcionario["usuario_id"],),
            )

        conn.commit()
        return jsonify({"status": "ok", "mensagem": "Vínculo desativado com sucesso."})

    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(exc)}), 400
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": "Não foi possível desativar o vínculo."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()






@views_bp.route("/atualizar_status_usuario", methods=["POST"])
@require_roles("administrador")
def atualizar_status_usuario():
    conn = None
    cursor = None
    try:
        dados = request.get_json(silent=True) or {}
        funcionario_id = uuid_obrigatorio(dados.get("funcionario_id"), "Funcionário")
        if str(funcionario_id) == str(session.get("funcionario_id")):
            return jsonify({"status": "erro", "mensagem": "Você não pode alterar seu próprio vínculo."}), 403

        conn = conectar_bd()
        cursor = conn.cursor(cursor_factory=psycopg2.extras.RealDictCursor)
        cursor.execute(
            "SELECT usuario_id, status FROM funcionarios WHERE id = %s AND empresa_id = %s FOR UPDATE",
            (funcionario_id, g.auth_context["empresa_id"]),
        )
        funcionario = cursor.fetchone()
        if not funcionario:
            conn.rollback()
            return jsonify({"status": "erro", "mensagem": "Funcionário não encontrado."}), 404

        novo_status = "desligado" if funcionario["status"] == "ativo" else "ativo"
        cursor.execute(
            "UPDATE funcionarios SET status = %s, updated_at = now() WHERE id = %s AND empresa_id = %s",
            (novo_status, funcionario_id, g.auth_context["empresa_id"]),
        )
        if novo_status == "ativo":
            cursor.execute(
                "UPDATE usuarios SET status = 'ativo', updated_at = now() WHERE id = %s",
                (funcionario["usuario_id"],),
            )
        else:
            cursor.execute(
                "SELECT 1 FROM funcionarios WHERE usuario_id = %s AND status = 'ativo' LIMIT 1",
                (funcionario["usuario_id"],),
            )
            if not cursor.fetchone():
                cursor.execute(
                    "UPDATE usuarios SET status = 'inativo', updated_at = now() WHERE id = %s",
                    (funcionario["usuario_id"],),
                )

        conn.commit()

        return jsonify({
            "status": "ok",
            "novo_status": novo_status
        })

    except ValueError as exc:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": str(exc)}), 400
    except Exception:
        if conn:
            conn.rollback()
        return jsonify({"status": "erro", "mensagem": "Não foi possível atualizar o status."}), 500
    finally:
        if cursor:
            cursor.close()
        if conn:
            conn.close()

# =========================
# STATUS (MOCK)
# =========================
contador = 0


@views_bp.route("/status")
def status():
    global contador

    estados = ["expediente", "intervalo", "fora"]
    estado = estados[contador % 3]

    contador += 1

    return {"status": estado}
